"""Generates the Phase 2 test corpus: realistic, deliberately NON-compliant originals (PDF, Word, HTML, text).

Run once (`python tests/corpus/build_corpus.py`); the outputs are committed. Needs reportlab and python-docx (dev only).
Content is a loose, restructured retelling of the ReadmeForge samples: different chapter order, prose instead of
tables, missing sections, company-specific wording. It is test input, never knowledge.
"""
import json
import shutil
from pathlib import Path

from docx import Document
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

HERE = Path(__file__).parent
INCOMING = HERE / "incoming"
REVISIONS = HERE / "revisions"

SRS_NAME = "ReadmeForge SRS v0.3.pdf"


def write_pdf(path: Path, title: str, sections: list[tuple[str, list[str]]]) -> None:
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for heading, paragraphs in sections:
        story.append(Paragraph(heading, styles["Heading2"]))
        for text in paragraphs:
            story.append(Paragraph(text, styles["BodyText"]))
            story.append(Spacer(1, 6))
    SimpleDocTemplate(str(path), title=title).build(story)


def srs(version: str, extra_requirement: bool) -> tuple[str, list[tuple[str, list[str]]]]:
    reqs = [
        "1.1 Whenever somebody pushes to the default branch of a linked repository we have to kick off a scan, "
        "and the user should see the outcome in the dashboard within about a minute. Without this the whole product is pointless.",
        "1.2 The tool must put together a README proposal from the scan and show it as a diff. Nothing gets committed "
        "unless the maintainer accepts it. (Nice for trust.)",
        "1.3 It would be good if maintainers could keep folders out of the scan via a .readmeforge-ignore file.",
    ]
    if extra_requirement:
        reqs.append("1.4 Maintainers can pause automatic scanning for a repository for up to 30 days. Added in this revision.")
    return (
        f"ReadmeForge - Requirements (working draft {version})",
        [
            ("0. About this paper", [
                "ReadmeForge is the product that writes and refreshes README files for GitHub repositories. "
                f"This is working draft {version}, written by the product owner for the platform team. Not yet reviewed.",
                "Two plans exist: Free (one repository, community support) and Pro (ten repositories, support desk).",
            ]),
            ("1. What it has to do", reqs),
            ("2. Quality expectations", [
                "2.1 The main page should load fast - we are aiming for under 2 seconds at the 95th percentile, checked monthly on staging.",
                "2.2 We need the dashboard and webhook receiver to be up 99.5% of the month during working hours.",
                "2.3 Only repositories the signed-in user authorised through GitHub OAuth may be read. Tested on every release.",
            ]),
            ("3. Things we will not do", [
                "No wikis, no API reference docs, no hosting providers other than GitHub, and never a commit or pull request "
                "without explicit approval.",
            ]),
        ],
    )


def build_srs() -> None:
    title, sections = srs("0.3", extra_requirement=False)
    write_pdf(INCOMING / SRS_NAME, title, sections)
    title, sections = srs("0.4", extra_requirement=True)
    REVISIONS.mkdir(exist_ok=True)
    write_pdf(REVISIONS / SRS_NAME, title, sections)
    shutil.copyfile(INCOMING / SRS_NAME, INCOMING / "RF-requirements-copy.pdf")  # same bytes, other name


def build_system_management_docx() -> None:
    doc = Document()
    doc.add_heading("ReadmeForge - How we run it", 0)
    doc.add_heading("Where it lives", 1)
    doc.add_paragraph(
        "A FastAPI backend and a React dashboard run as Azure Container Apps in four Indian Azure locations: "
        "Central, South, West and Jio India West. Data sits in Azure Database for PostgreSQL, messages go through "
        "Azure Service Bus, Redis caches scan state and Blob Storage keeps scan artefacts."
    )
    doc.add_heading("Environments", 1)
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ["Env", "Purpose", "Notes"]):
        cell.text = text
    for row in [("dev", "Developer testing", "Single replica, shared Postgres"),
                ("staging", "Pre-release checks and monthly load test", "Mirrors prod sizing at 25%"),
                ("prod", "Customers", "Four locations behind Azure Front Door")]:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    doc.add_heading("Shipping a release", 1)
    doc.add_paragraph("Merge to main, the pipeline builds one image, deploys to staging, waits for smoke tests, then "
                      "a human approves the production rollout. Rollback is traffic shifting to the previous revision.")
    doc.add_heading("Watching it", 1)
    doc.add_paragraph("Datadog monitors API 5xx rate, queue age and PostgreSQL CPU. P1 alerts page the on-call through PagerDuty.")
    doc.add_heading("Backups", 1)
    doc.add_paragraph("PostgreSQL point-in-time restore for 14 days; Blob Storage soft delete for 30 days. Restore was last tested in Q2.")
    doc.add_heading("Who helps", 1)
    doc.add_paragraph("Free users: community forum only. Pro users: support desk, first response within one business day.")
    # No SOPs, no known errors, no SLOs: this document is intentionally incomplete.
    doc.save(INCOMING / "ReadmeForge_SystemMgmt.docx")


def build_runbooks_html() -> None:
    html = """<!doctype html>
<html><head><meta charset="utf-8"><title>ReadmeForge on-call runbooks</title></head>
<body>
<h1>ReadmeForge on-call runbooks</h1>
<p>Last edited by the platform team. Everything below applies to <b>production</b> only.</p>
<h2>API 5xx spike (monitor: [ReadmeForge][prod] API 5xx error rate high)</h2>
<ol>
<li>Acknowledge the PagerDuty incident and say "Investigating" in #readmeforge-alerts.</li>
<li>Open the dashboard "ReadmeForge - API overview" and find which location carries the errors.</li>
<li>If errors started within 30 minutes of a deployment, roll the api app back with
<code>az containerapp ingress traffic set</code> to the previous revision.</li>
<li>Done when the 5xx ratio stays below 1% for 15 minutes. Escalate to the Platform Engineering Lead after 30 minutes.</li>
</ol>
<h2>PostgreSQL CPU saturation</h2>
<ol>
<li>Check Datadog "PostgreSQL and Redis health" and find the heaviest queries.</li>
<li>Scale the server up one tier if CPU stays above 90% for 15 minutes.</li>
<li>Tell the on-call L3 engineer, who decides on a permanent fix.</li>
</ol>
</body></html>
"""
    (INCOMING / "readmeforge-runbooks.html").write_text(html, encoding="utf-8")


def build_shared_standards_pdf() -> None:
    write_pdf(INCOMING / "Platform Reliability Standards 2026.pdf", "Platform Reliability Standards 2026", [
        ("Scope", ["Applies to every production application on the platform unless a service owner agrees a stricter target."]),
        ("Targets", [
            "Production availability: 99.9% per calendar month, measured as non-5xx responses to valid requests.",
            "API latency: 95th percentile under 800 ms over a rolling 7 days.",
            "Priority 1 incidents: acknowledged within 30 minutes (24x7) and restored within 4 hours.",
            "Free or community tiers carry no contractual commitment.",
        ]),
        ("Review", ["Service owners and the platform management team review these targets once a year."]),
    ])


def build_ambiguous_notes() -> None:
    (INCOMING / "notes.txt").write_text(
        "call back re: the thing from tuesday\n- check numbers w/ finance\n- maybe move the offsite?\n"
        "- ask Priya about the template\n",
        encoding="utf-8",
    )


EXPECTATIONS = {
    SRS_NAME: {"scope": "ReadmeForge", "outcome": "stored"},
    "RF-requirements-copy.pdf": {"scope": "ReadmeForge", "outcome": "stored or SAME-as-existing", "note": "same bytes as the SRS under another name: identity is folder+name so it is NEW, an agent may mention the duplicate"},
    "ReadmeForge_SystemMgmt.docx": {"scope": "ReadmeForge", "outcome": "stored"},
    "readmeforge-runbooks.html": {"scope": "ReadmeForge", "outcome": "stored"},
    "Platform Reliability Standards 2026.pdf": {"scope": "shared", "outcome": "stored"},
    "notes.txt": {"scope": None, "outcome": "deferred"},
    "revisions/" + SRS_NAME: {"scope": "ReadmeForge", "outcome": "stored as CHANGED (version 2) when copied into incoming/ after the first run"},
}


def main() -> None:
    INCOMING.mkdir(parents=True, exist_ok=True)
    build_srs()
    build_system_management_docx()
    build_runbooks_html()
    build_shared_standards_pdf()
    build_ambiguous_notes()
    (HERE / "corpus_expectations.json").write_text(json.dumps(EXPECTATIONS, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
